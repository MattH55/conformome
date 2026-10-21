#!/usr/bin/env python3
"""Post-process the pandoc docx: three-rule tables (booktabs look), 10 pt table text,
bold header row, page numbers in the footer."""
import sys

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

path = sys.argv[1]
d = docx.Document(path)


def borders(el, tag, spec):
    pr = el.find(qn(tag))
    if pr is None:
        pr = OxmlElement(tag)
        el.insert(0, pr)
    b = pr.find(qn("w:tblBorders" if tag == "w:tblPr" else "w:tcBorders"))
    if b is not None:
        pr.remove(b)
    b = OxmlElement("w:tblBorders" if tag == "w:tblPr" else "w:tcBorders")
    for side, (val, sz) in spec.items():
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), val)
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "000000")
        b.append(e)
    pr.append(b)


for t in d.tables:
    tblPr = t._tbl.tblPr
    borders(t._tbl, "w:tblPr", {"top": ("single", 12), "bottom": ("single", 12), "left": ("nil", 0),
                                "right": ("nil", 0), "insideH": ("nil", 0), "insideV": ("nil", 0)})
    for i, row in enumerate(t.rows):
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            if i == 0:
                borders(cell._tc, "w:tcPr", {"bottom": ("single", 6)})
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    r.font.size = Pt(10)
                    if i == 0:
                        r.font.bold = True

# page number field, centred in footer
for sec in d.sections:
    fp = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
    fp.alignment = 1
    run = fp.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), kind); run._r.append(f)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; run._r.append(it)
    run.font.size = Pt(10)

d.save(path)
print("polished", path)
